import logging
from pathlib import Path
from typing import Any

import chardet
import openpyxl
import pandas as pd

from app.core.config import settings

logger = logging.getLogger(__name__)


class FileProcessor:
    """Excel/CSVファイルの処理クラス"""

    @staticmethod
    def _detect_encoding(file_path: Path) -> str:
        """
        ファイルのエンコーディングを検出

        Args:
            file_path: ファイルパス

        Returns:
            検出されたエンコーディング（デフォルト: utf-8）
        """
        try:
            with open(file_path, "rb") as f:
                raw_data = f.read(10000)  # 最初の10KBを読み込み
                detected = chardet.detect(raw_data)
                encoding = detected["encoding"] or "utf-8"
                logger.debug(
                    f"Detected encoding for {file_path}: {encoding} (confidence: {detected.get('confidence', 0)})"
                )
                return encoding
        except Exception as e:
            logger.warning(
                f"Failed to detect encoding for {file_path}: {e}. Using utf-8 as fallback."
            )
            return "utf-8"

    @staticmethod
    def extract_metadata(file_path: Path, file_type: str) -> dict[str, Any]:
        """
        ファイルからメタデータを抽出

        Args:
            file_path: ファイルパス
            file_type: ファイルタイプ (.xlsx, .xls, .csv)

        Returns:
            メタデータ辞書 (rows_count, columns_count, sheets_count)
        """
        metadata = {
            "rows_count": None,
            "columns_count": None,
            "sheets_count": None,
        }

        try:
            if file_type in [".xlsx", ".xls"]:
                metadata = FileProcessor._extract_excel_metadata(file_path)
            elif file_type == ".csv":
                metadata = FileProcessor._extract_csv_metadata(file_path)
        except Exception as e:
            logger.error(f"Error extracting metadata from {file_path}: {e}", exc_info=True)

        return metadata

    @staticmethod
    def _extract_excel_metadata(file_path: Path) -> dict[str, Any]:
        """Excelファイルからメタデータを抽出"""
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        sheets_count = len(wb.sheetnames)

        # 最初のシートから行数・列数を取得
        first_sheet = wb[wb.sheetnames[0]]
        rows_count = first_sheet.max_row
        columns_count = first_sheet.max_column

        wb.close()

        return {
            "rows_count": rows_count,
            "columns_count": columns_count,
            "sheets_count": sheets_count,
        }

    @staticmethod
    def _extract_csv_metadata(file_path: Path) -> dict[str, Any]:
        """CSVファイルからメタデータを抽出"""
        # エンコーディングを検出
        encoding = FileProcessor._detect_encoding(file_path)

        df = pd.read_csv(file_path, nrows=0, encoding=encoding)  # ヘッダーのみ読み込み
        columns_count = len(df.columns)

        # 行数をカウント（効率的に）
        with open(file_path, encoding=encoding) as f:
            rows_count = sum(1 for _ in f) - 1  # ヘッダーを除く

        return {
            "rows_count": rows_count,
            "columns_count": columns_count,
            "sheets_count": 1,  # CSVは常に1シート
        }

    @staticmethod
    def get_preview_data(
        file_path: Path, file_type: str, max_rows: int = None
    ) -> tuple[list[str], list[list[Any]]]:
        """
        プレビュー用のデータを取得

        Args:
            file_path: ファイルパス
            file_type: ファイルタイプ
            max_rows: 最大行数（Noneの場合は設定値を使用）

        Returns:
            (ヘッダーリスト, データ行のリスト)
        """
        if max_rows is None:
            max_rows = settings.MAX_ROWS_PREVIEW

        try:
            if file_type in [".xlsx", ".xls"]:
                return FileProcessor._get_excel_preview(file_path, max_rows)
            elif file_type == ".csv":
                return FileProcessor._get_csv_preview(file_path, max_rows)
            else:
                logger.warning(f"Unsupported file type: {file_type}")
                return [], []
        except Exception as e:
            logger.error(f"Error getting preview data from {file_path}: {e}", exc_info=True)
            return [], []

    @staticmethod
    def _get_excel_preview(file_path: Path, max_rows: int) -> tuple[list[str], list[list[Any]]]:
        """Excelファイルのプレビューデータを取得"""
        df = pd.read_excel(file_path, nrows=max_rows)
        headers = df.columns.tolist()
        rows = df.values.tolist()

        # NaN を None に変換
        rows = [[None if pd.isna(cell) else cell for cell in row] for row in rows]

        return headers, rows

    @staticmethod
    def _get_csv_preview(file_path: Path, max_rows: int) -> tuple[list[str], list[list[Any]]]:
        """CSVファイルのプレビューデータを取得"""
        # エンコーディングを検出
        encoding = FileProcessor._detect_encoding(file_path)

        df = pd.read_csv(file_path, nrows=max_rows, encoding=encoding)
        headers = df.columns.tolist()
        rows = df.values.tolist()

        # NaN を None に変換
        rows = [[None if pd.isna(cell) else cell for cell in row] for row in rows]

        return headers, rows

    @staticmethod
    def convert_to_dict(headers: list[str], rows: list[list[Any]]) -> list[dict[str, Any]]:
        """ヘッダーと行データを辞書のリストに変換"""
        return [dict(zip(headers, row, strict=False)) for row in rows]
